"""M2 客户接口：CRUD + 去重 + 合并 + 编码。"""

from __future__ import annotations

import io
import uuid
from typing import Any

from fastapi import APIRouter, Depends, File, Query, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, or_, select, update

from app.core import response as resp
from app.core.deps import DbSession, require_permission
from app.core.error_codes import ErrCode
from app.core.exceptions import ApiError, ConflictError, ForbiddenError
from app.models.business import (
    Activity,
    BusinessCode,
    Company,
    Contact,
    Inquiry,
    Opportunity,
)
from app.models.email import EmailThread
from app.models.pricing import Quotation
from app.models.user import User
from app.schemas.business import (
    BusinessCodeOut,
    CompanyBatchDeleteIn,
    CompanyBatchUpdateIn,
    CompanyCreate,
    CompanyMergeIn,
    CompanyMergeOut,
    CompanyOut,
    CompanyUpdate,
    CurrencyCodeIn,
    CurrencyCodeOut,
    CustomerCodePreview,
)
from app.services import company_code, ebms_sync
from app.services.data_scope import apply_data_scope, can_access_owner, get_user_scope
from app.services.query import apply_sort, get_or_404, paginate, soft_delete

router = APIRouter(prefix="/customers", tags=["M2-客户"])

CODE_RULE = "{业务员编号}-{业务员第N个客户}-{公司第N个客户}，如 HJ003-176-576；同客户多币别追加 -1/-2"

#: 允许批量修改的客户字段白名单（对齐 EBMS 批量编辑，仅覆盖表单标量字段）
COMPANY_BATCH_EDITABLE_FIELDS = {
    "customer_level",
    "sales_mode",
    "customer_status",
    "status",
    "country",
    "country_region",
    "industry",
    "source",
}

#: Excel 导入表头映射：EBMS 模板中文列名 -> CRM Company 字段（客户编号由系统自动生成，不参与映射）
CUSTOMER_IMPORT_HEADER_MAP = {
    "客户全称": "name",
    "客户曾用名": "former_name",
    "母公司编号": "parent_code",
    "决策人姓名": "decision_maker",
    "决策人职位": "decision_maker_position",
    "决策核心诉求": "decision_core_need",
    "客户等级": "customer_level",
    "销售模式": "sales_mode",
    "客户状态": "customer_status",
    "最后交易年份": "last_trade_year",
    "开票抬头": "invoice_title",
    "纳税人识别号": "tax_id",
    "开户银行": "bank_name",
    "银行账号": "bank_account",
    "地址": "address",
    "国家/地区": "country_region",
    "客户来源": "source",
    "备注说明": "remark",
    "网站": "website",
    "行业": "industry",
}


async def _company_or_404(db: DbSession, company_id: uuid.UUID) -> Company:
    return await get_or_404(db, Company, company_id, "客户")


async def _enrich(db: DbSession, rows: list[Company]) -> list[dict[str, Any]]:
    """补全负责人名称、联系人数量、编码规则提示。"""
    ids = [r.id for r in rows]
    counts: dict[uuid.UUID, int] = {}
    if ids:
        found = (
            await db.execute(
                select(Contact.company_id, func.count())
                .where(Contact.company_id.in_(ids), Contact.deleted_at.is_(None))
                .group_by(Contact.company_id)
            )
        ).all()
        counts = {r[0]: int(r[1]) for r in found}

    owner_ids = {r.owner_id for r in rows if r.owner_id}
    owners: dict[uuid.UUID, str] = {}
    if owner_ids:
        found_owners = (await db.execute(select(User.id, User.name).where(User.id.in_(owner_ids)))).all()
        owners = {r[0]: r[1] for r in found_owners}

    result = []
    for row in rows:
        data = CompanyOut.model_validate(row).model_dump(mode="json")
        data["contact_count"] = counts.get(row.id, 0)
        data["owner_name"] = owners.get(row.owner_id) if row.owner_id else None
        data["code_rule"] = CODE_RULE
        result.append(data)
    return result


@router.get("", summary="客户列表")
async def list_companies(
    db: DbSession,
    current: User = Depends(require_permission("customer:view")),
    keyword: str | None = Query(None, description="名称/编码/国家"),
    country: str | None = None,
    customer_level: str | None = None,
    status: str | None = None,
    owner_id: uuid.UUID | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    sort: str = "created_at",
    order: str = Query("desc", pattern="^(asc|desc)$"),
) -> dict[str, Any]:
    stmt = select(Company).where(Company.deleted_at.is_(None))
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(
            or_(
                Company.name.ilike(like),
                Company.name_en.ilike(like),
                Company.code.ilike(like),
                Company.country.ilike(like),
            )
        )
    if country:
        stmt = stmt.where(Company.country == country)
    if customer_level:
        stmt = stmt.where(Company.customer_level == customer_level)
    if status:
        stmt = stmt.where(Company.status == status)
    if owner_id:
        stmt = stmt.where(Company.owner_id == owner_id)
    stmt = await apply_data_scope(stmt, Company, current, db)

    stmt = apply_sort(stmt, Company, sort, order)
    rows, total = await paginate(db, stmt, page, page_size)
    items = await _enrich(db, list(rows))
    return resp.page(items, total, page, page_size)


@router.post("", summary="新建客户", status_code=201)
async def create_company(
    payload: CompanyCreate,
    db: DbSession,
    current: User = Depends(require_permission("customer:create")),
) -> dict[str, Any]:
    owner_id = payload.owner_id or current.id
    owner = (await db.execute(select(User).where(User.id == owner_id))).scalar_one_or_none() or current

    code, owner_seq, global_seq = await company_code.next_company_code(owner)
    company = Company(
        **payload.model_dump(exclude_unset=True),
        code=code,
        owner_seq=owner_seq,
        global_seq=global_seq,
        owner_id=owner.id,
        department_id=current.department_id,
    )
    db.add(company)
    await db.flush()
    await company_code.ensure_primary_code(db, company)
    if company.currency:
        await company_code.add_currency_code(db, company, company.currency)
    await ebms_sync.emit_upsert(db, "customer", company)
    await db.commit()
    await db.refresh(company)
    items = await _enrich(db, [company])
    return resp.ok(items[0])


@router.get("/next-code", summary="预览下一个客户编码")
async def next_code_preview(
    current: User = Depends(require_permission("customer:view")),
) -> dict[str, Any]:
    code, owner_seq, global_seq = await company_code.preview_next_code(current)
    return resp.ok(
        CustomerCodePreview(
            code=code,
            owner_code=current.code,
            owner_seq=owner_seq,
            global_seq=global_seq,
            rule=CODE_RULE,
        ).model_dump(mode="json")
    )


@router.get("/duplicates", summary="客户重复检测")
async def duplicates(
    db: DbSession,
    _: User = Depends(require_permission("customer:view")),
) -> dict[str, Any]:
    groups = (
        await db.execute(
            select(Company.name, func.count().label("cnt"))
            .where(Company.deleted_at.is_(None))
            .group_by(Company.name)
            .having(func.count() > 1)
        )
    ).all()
    result = []
    for name, _cnt in groups:
        rows = (
            (await db.execute(select(Company).where(Company.name == name, Company.deleted_at.is_(None))))
            .scalars()
            .all()
        )
        result.append(
            {
                "key": name,
                "companies": [CompanyOut.model_validate(r).model_dump(mode="json") for r in rows],
            }
        )
    return resp.ok(result)


@router.post("/merge", summary="合并客户")
async def merge_companies(
    payload: CompanyMergeIn,
    db: DbSession,
    _: User = Depends(require_permission("customer:merge")),
) -> dict[str, Any]:
    target = await _company_or_404(db, payload.keep_id)
    source_ids = [sid for sid in payload.source_ids if sid != payload.keep_id]
    if not source_ids:
        raise ConflictError("未指定需要合并的客户")

    moved_contacts = moved_opportunities = 0
    for source_id in source_ids:
        await _company_or_404(db, source_id)

    for source_id in source_ids:
        result = await db.execute(update(Contact).where(Contact.company_id == source_id).values(company_id=target.id))
        moved_contacts += int(getattr(result, "rowcount", 0) or 0)
        result = await db.execute(
            update(Opportunity).where(Opportunity.company_id == source_id).values(company_id=target.id)
        )
        moved_opportunities += int(getattr(result, "rowcount", 0) or 0)
        await db.execute(update(Inquiry).where(Inquiry.company_id == source_id).values(company_id=target.id))
        await db.execute(update(Quotation).where(Quotation.company_id == source_id).values(company_id=target.id))
        await db.execute(update(Activity).where(Activity.company_id == source_id).values(company_id=target.id))
        await db.execute(update(EmailThread).where(EmailThread.company_id == source_id).values(company_id=target.id))
        await db.execute(update(BusinessCode).where(BusinessCode.company_id == source_id).values(company_id=target.id))
        source = await _company_or_404(db, source_id)
        source.status = "MERGED"
        source.remark = f"{source.remark or ''}\n已合并至 {target.code}".strip()
        soft_delete(source)
        await ebms_sync.emit_delete(db, "customer", source)

    await ebms_sync.emit_upsert(db, "customer", target)
    await db.commit()
    return resp.ok(
        CompanyMergeOut(
            target_id=target.id,
            merged=len(source_ids),
            moved_contacts=moved_contacts,
            moved_opportunities=moved_opportunities,
        ).model_dump(mode="json")
    )


@router.put("/batch-update", summary="批量修改客户（仅白名单字段）")
async def batch_update_companies(
    payload: CompanyBatchUpdateIn,
    db: DbSession,
    current: User = Depends(require_permission("customer:update")),
) -> dict[str, Any]:
    unknown = set(payload.fields) - COMPANY_BATCH_EDITABLE_FIELDS
    if unknown:
        raise ApiError(
            ErrCode.VALIDATION_ERROR,
            f"不允许批量修改字段: {', '.join(sorted(unknown))}",
            http_status=400,
        )
    stmt = select(Company).where(Company.id.in_(payload.ids), Company.deleted_at.is_(None))
    stmt = await apply_data_scope(stmt, Company, current, db)
    rows = (await db.execute(stmt)).scalars().all()
    updated = 0
    for row in rows:
        for key, value in payload.fields.items():
            setattr(row, key, value)
        await ebms_sync.emit_upsert(db, "customer", row)
        updated += 1
    await db.commit()
    return resp.ok({"updated": updated, "requested": len(payload.ids)})


@router.post("/batch-delete", summary="批量删除客户（软删除）")
async def batch_delete_companies(
    payload: CompanyBatchDeleteIn,
    db: DbSession,
    current: User = Depends(require_permission("customer:delete")),
) -> dict[str, Any]:
    stmt = select(Company).where(Company.id.in_(payload.ids), Company.deleted_at.is_(None))
    stmt = await apply_data_scope(stmt, Company, current, db)
    rows = (await db.execute(stmt)).scalars().all()
    deleted = 0
    for row in rows:
        soft_delete(row)
        await ebms_sync.emit_delete(db, "customer", row)
        deleted += 1
    await db.commit()
    return resp.ok({"deleted": deleted, "requested": len(payload.ids)})


@router.post("/import-xlsx", summary="Excel 批量导入客户（按名称去重，同一文件重复导入不产生重复）")
async def import_companies_xlsx(
    db: DbSession,
    file: UploadFile = File(...),
    current: User = Depends(require_permission("customer:create")),
) -> dict[str, Any]:
    content = await file.read()
    if not content:
        raise ApiError(ErrCode.VALIDATION_ERROR, "请上传非空文件", http_status=400)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:  # noqa: BLE001
        raise ApiError(ErrCode.VALIDATION_ERROR, f"无法解析 Excel 文件: {exc}", http_status=400) from exc

    if not rows:
        return resp.ok({"imported": 0, "skipped": 0, "results": []})

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    col_map = {idx: CUSTOMER_IMPORT_HEADER_MAP[h] for idx, h in enumerate(headers) if h in CUSTOMER_IMPORT_HEADER_MAP}

    existing_names = {
        name.strip().lower()
        for (name,) in (
            await db.execute(select(Company.name).where(Company.deleted_at.is_(None)))
        ).all()
        if name
    }

    imported = 0
    skipped = 0
    results: list[dict[str, Any]] = []
    seen_names: set[str] = set()
    for line_no, row in enumerate(rows[1:], start=2):
        if not row:
            continue
        fields: dict[str, Any] = {}
        for idx, field in col_map.items():
            cell = row[idx] if idx < len(row) else None
            if cell is not None:
                fields[field] = str(cell).strip()
        name = fields.get("name") or ""
        if not name:
            skipped += 1
            results.append({"row": line_no, "status": "skipped", "reason": "客户名称为空"})
            continue
        normalized = name.strip().lower()
        if normalized in existing_names or normalized in seen_names:
            skipped += 1
            results.append({"row": line_no, "name": name, "status": "skipped", "reason": "客户名称已存在"})
            continue

        code, owner_seq, global_seq = await company_code.next_company_code(current)
        company = Company(
            **fields,
            code=code,
            owner_seq=owner_seq,
            global_seq=global_seq,
            owner_id=current.id,
            department_id=current.department_id,
            currency="USD",
        )
        db.add(company)
        await db.flush()
        await company_code.ensure_primary_code(db, company)
        await ebms_sync.emit_upsert(db, "customer", company)
        seen_names.add(normalized)
        imported += 1
        results.append({"row": line_no, "name": name, "status": "imported"})

    await db.commit()
    return resp.ok({"imported": imported, "skipped": skipped, "results": results})


@router.get("/{company_id}", summary="客户详情")
async def get_company(
    company_id: uuid.UUID,
    db: DbSession,
    current: User = Depends(require_permission("customer:view")),
) -> dict[str, Any]:
    company = await _company_or_404(db, company_id)
    scope_type, dept_ids = await get_user_scope(current, db)
    if not can_access_owner(current, scope_type, dept_ids, company.owner_id, company.department_id):
        raise ForbiddenError("无权限访问该客户")
    enriched = await _enrich(db, [company])
    data = enriched[0]

    codes = await company_code.codes_of(db, company_id)
    data["business_codes"] = [BusinessCodeOut.model_validate(c).model_dump(mode="json") for c in codes]

    contacts = (
        (
            await db.execute(
                select(Contact)
                .where(Contact.company_id == company_id, Contact.deleted_at.is_(None))
                .order_by(Contact.is_primary.desc(), Contact.created_at.asc())
            )
        )
        .scalars()
        .all()
    )
    data["contacts"] = [
        {
            "id": str(c.id),
            "name": c.name,
            "position": c.position,
            "email": c.email,
            "phone": c.phone,
            "mobile": c.mobile,
            "is_primary": c.is_primary,
        }
        for c in contacts
    ]

    opportunities = (
        (
            await db.execute(
                select(Opportunity)
                .where(Opportunity.company_id == company_id, Opportunity.deleted_at.is_(None))
                .order_by(Opportunity.created_at.desc())
                .limit(50)
            )
        )
        .scalars()
        .all()
    )
    data["opportunities"] = [
        {
            "id": str(o.id),
            "opportunity_no": o.opportunity_no,
            "name": o.name,
            "stage": o.stage,
            "estimated_amount": float(o.estimated_amount) if o.estimated_amount else None,
            "currency": o.currency,
        }
        for o in opportunities
    ]

    inquiries = (
        (
            await db.execute(
                select(Inquiry)
                .where(Inquiry.company_id == company_id, Inquiry.deleted_at.is_(None))
                .order_by(Inquiry.created_at.desc())
                .limit(50)
            )
        )
        .scalars()
        .all()
    )
    data["inquiries"] = [
        {
            "id": str(i.id),
            "inquiry_no": i.inquiry_no,
            "status": i.status,
            "currency": i.currency,
            "created_at": i.created_at.isoformat() if i.created_at else None,
        }
        for i in inquiries
    ]

    emails = (
        (
            await db.execute(
                select(EmailThread)
                .where(EmailThread.company_id == company_id, EmailThread.deleted_at.is_(None))
                .order_by(EmailThread.last_message_at.desc().nullslast())
                .limit(20)
            )
        )
        .scalars()
        .all()
    )
    data["email_threads"] = [
        {
            "id": str(t.id),
            "subject": t.subject,
            "message_count": t.message_count,
            "last_message_at": t.last_message_at.isoformat() if t.last_message_at else None,
        }
        for t in emails
    ]
    return resp.ok(data)


@router.put("/{company_id}", summary="更新客户")
async def update_company(
    company_id: uuid.UUID,
    payload: CompanyUpdate,
    db: DbSession,
    _: User = Depends(require_permission("customer:update")),
) -> dict[str, Any]:
    company = await _company_or_404(db, company_id)
    data = payload.model_dump(exclude_unset=True)
    new_currency = data.get("currency")
    for key, value in data.items():
        setattr(company, key, value)
    if new_currency:
        await company_code.add_currency_code(db, company, new_currency)
    await ebms_sync.emit_upsert(db, "customer", company)
    await db.commit()
    await db.refresh(company)
    items = await _enrich(db, [company])
    return resp.ok(items[0])


@router.delete("/{company_id}", summary="删除客户（软删除）")
async def delete_company(
    company_id: uuid.UUID,
    db: DbSession,
    _: User = Depends(require_permission("customer:delete")),
) -> dict[str, Any]:
    company = await _company_or_404(db, company_id)
    soft_delete(company)
    await ebms_sync.emit_delete(db, "customer", company)
    await db.commit()
    return resp.ok({"deleted": True})


@router.get("/{company_id}/business-codes", summary="客户业务编码（含币别码）")
async def list_business_codes(
    company_id: uuid.UUID,
    db: DbSession,
    _: User = Depends(require_permission("customer:view")),
) -> dict[str, Any]:
    await _company_or_404(db, company_id)
    rows = await company_code.codes_of(db, company_id)
    return resp.ok([BusinessCodeOut.model_validate(r).model_dump(mode="json") for r in rows])


@router.post("/{company_id}/currencies", summary="为客户追加币别编码（-1/-2…）", status_code=201)
async def add_currency_code(
    company_id: uuid.UUID,
    payload: CurrencyCodeIn,
    db: DbSession,
    _: User = Depends(require_permission("customer:update")),
) -> dict[str, Any]:
    company = await _company_or_404(db, company_id)
    row, created = await company_code.add_currency_code(db, company, payload.currency)
    await db.commit()
    await db.refresh(row)
    return resp.ok(
        CurrencyCodeOut(id=row.id, code=row.code, currency=row.currency, created=created).model_dump(mode="json")
    )
