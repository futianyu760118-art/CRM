"""Excel 导出：询价单（客户询价单模板）/ 报价单（对外价格表模板）。

版式对齐用户提供的两份实际单据：
- 询价单：《销2询价.xlsx》一个页签 = 一张询价
- 报价单：《星耀内置混动.xlsx》英文产品价格表
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.business import Company, Inquiry, InquiryItem
from app.models.email import File
from app.models.pricing import Quotation, QuotationItem
from app.services import storage

TITLE_FONT = Font(bold=True, size=14)
HEAD_FONT = Font(bold=True, size=10)
NORMAL_FONT = Font(size=10)
NOTE_FONT = Font(size=9)
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREY = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


async def _images(db: AsyncSession, file_ids: list[Any]) -> dict[str, bytes]:
    ids = [f for f in file_ids if f]
    if not ids:
        return {}
    rows = (await db.execute(select(File).where(File.id.in_(ids)))).scalars().all()
    result: dict[str, bytes] = {}
    for row in rows:
        try:
            result[str(row.id)] = storage.read_bytes(row)
        except Exception:  # noqa: BLE001
            continue
    return result


def _insert_image(ws, data: bytes, anchor: str, width: int = 90, height: int = 70) -> None:
    try:
        img = XLImage(io.BytesIO(data))
        ratio = (img.height / img.width) if img.width else 1
        img.width = width
        img.height = int(width * ratio) if ratio else height
        ws.add_image(img, anchor)
    except Exception:  # noqa: BLE001
        pass


# ------------------------------------------------------------------ 询价单


async def build_inquiry_workbook(
    db: AsyncSession, inquiry: Inquiry, *, company_title: str, version_default: str = "2025-08"
) -> bytes:
    items = list(
        (
            await db.execute(
                select(InquiryItem)
                .where(InquiryItem.inquiry_id == inquiry.id, InquiryItem.deleted_at.is_(None))
                .order_by(InquiryItem.sort, InquiryItem.created_at)
            )
        )
        .scalars()
        .all()
    )
    images = await _images(db, [i.image_file_id for i in items])

    wb = Workbook()
    ws = wb.active
    ws.title = (inquiry.inquiry_no or "INQUIRY")[-12:]

    ws.merge_cells("A1:J1")
    ws["A1"] = company_title
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("Q1:V1")
    ws["Q1"] = "客户产品图片"
    ws["Q1"].alignment = CENTER

    ws.merge_cells("A2:J2")
    ws["A2"] = f"版本号：{inquiry.version_no or version_default}"
    ws["A2"].font = NOTE_FONT

    ws.merge_cells("A3:J3")
    ws["A3"] = "客户询价单"
    ws["A3"].font = TITLE_FONT
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A4:V4")
    ws["A4"] = "询价表操作标准：和已经有过类似或者标准报价有差异点进行说明，表上信息填写齐全"
    ws["A4"].font = NOTE_FONT

    head_labels = [
        "名称",
        "部门",
        "业务员",
        "客户编号",
        "客户名",
        "新/老客户",
        "询价单号",
        "发起时间",
        "结束时间",
        "样品可供时间",
        "下单时间",
        "出货时间",
        "首单货值（万）",
        "数量（个)",
        "年金额",
        "市场",
    ]
    head_values: list[Any] = [
        inquiry.customer_name or "",
        inquiry.department_name_text or "",
        inquiry.owner_name_text or "",
        inquiry.customer_code or "",
        inquiry.customer_name or "",
        {"NEW": "新", "OLD": "老"}.get(inquiry.customer_type or "", inquiry.customer_type or ""),
        inquiry.inquiry_no,
        _fmt_date(inquiry.start_date),
        _fmt_date(inquiry.end_date),
        _fmt_date(inquiry.sample_date),
        _fmt_date(inquiry.order_date),
        _fmt_date(inquiry.shipment_date),
        inquiry.first_order_amount,
        inquiry.quantity_total,
        inquiry.annual_amount,
        inquiry.market or "",
    ]
    for idx, label in enumerate(head_labels, start=1):
        cell = ws.cell(row=5, column=idx, value=label)
        cell.font = HEAD_FONT
        cell.fill = GREY
        cell.border = BORDER
        cell.alignment = CENTER
        value_cell = ws.cell(row=6, column=idx, value=head_values[idx - 1])
        value_cell.font = NORMAL_FONT
        value_cell.border = BORDER
        value_cell.alignment = WRAP

    ws.cell(row=8, column=1, value="交样周期").font = HEAD_FONT
    ws.cell(row=8, column=2, value=inquiry.sample_lead_time or "")
    ws.cell(row=8, column=4, value="已有产品").font = HEAD_FONT
    ws.cell(row=8, column=8, value="需开发产品").font = HEAD_FONT

    new_items = [i for i in items if (i.group_type or "EXISTING") == "NEW"]
    existing_items = [i for i in items if (i.group_type or "EXISTING") != "NEW"]

    row = 10
    row = _write_standard_block(ws, row, inquiry.spec_rows or [])
    row += 1
    if new_items:
        row = _write_inquiry_block(ws, row, new_items, images, block="NEW")
        row += 1
    if existing_items:
        row = _write_inquiry_block(ws, row, existing_items, images, block="EXISTING")

    widths = [10, 8, 9, 9, 12, 8, 16, 11, 10, 12, 10, 10, 12, 10, 10, 10, 16, 12, 10, 12, 10, 22]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_standard_block(ws, start_row: int, spec_rows: list[dict[str, Any]]) -> int:
    """写出「询价单填写模板（客户标准）」纵向块：字段为行、产品为列。"""
    if not spec_rows:
        return start_row
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=5)
    title = ws.cell(row=start_row, column=1, value="询价单填写模板（客户标准）")
    title.font = HEAD_FONT
    ws.cell(row=start_row, column=7, value="客户产品图片").font = NOTE_FONT
    current = start_row + 1
    for spec in spec_rows:
        label = str(spec.get("label") or "")
        values = spec.get("values") or []
        label_cell = ws.cell(row=current, column=1, value=label)
        label_cell.font = HEAD_FONT
        label_cell.fill = GREY
        label_cell.border = BORDER
        label_cell.alignment = CENTER
        for idx, value in enumerate(values, start=2):
            cell = ws.cell(row=current, column=idx, value=value)
            cell.font = NORMAL_FONT
            cell.border = BORDER
            cell.alignment = WRAP
        current += 1
    return current


def _write_inquiry_block(ws, start_row: int, items: list[InquiryItem], images: dict[str, bytes], *, block: str) -> int:
    if block == "NEW":
        labels = [
            ("序号", "sort"),
            ("产品图", "_image"),
            ("产品型号", "product_model"),
            ("产品名称", "product_name"),
            ("功率", "power"),
            ("输入电压", "voltage"),
            ("电池", "battery"),
            ("色温", "color_temp"),
            ("光通量/光效", "lumen"),
            ("光源", "light_source"),
            ("主体", "main_body"),
            ("灯罩", "lamp_shade"),
            ("反光罩", "reflector"),
            ("电缆线", "cable"),
            ("开关", "switch"),
            ("USB", "usb"),
            ("防水等级", "waterproof"),
            ("感应器", "sensor"),
            ("报价", "quote_price"),
            ("目标价", "target_price"),
            ("需求数量", "quantity"),
            ("报价需时间", "quote_lead_time"),
            ("其他要求", "other_requirement"),
        ]
    else:
        labels = [
            ("", "sort"),
            ("产品型号", "product_model"),
            ("产品名称", "product_name"),
            ("功率", "power"),
            ("输入电压", "voltage"),
            ("电池", "battery"),
            ("色温", "color_temp"),
            ("光通量/光效", "lumen"),
            ("光源", "light_source"),
            ("主体", "main_body"),
            ("压框", "frame"),
            ("灯罩", "lamp_shade"),
            ("反光罩", "reflector"),
            ("电缆线", "cable"),
            ("开关", "switch"),
            ("USB", "usb"),
            ("防水等级", "waterproof"),
            ("感应器", "sensor"),
            ("其他要求1", "other_requirement"),
            ("其他要求2", "other_requirement2"),
            ("报价", "quote_price"),
        ]

    for idx, (label, _field) in enumerate(labels, start=1):
        cell = ws.cell(row=start_row, column=idx, value=label)
        cell.font = HEAD_FONT
        cell.fill = GREY
        cell.border = BORDER
        cell.alignment = CENTER

    ws.row_dimensions[start_row].height = 22
    current = start_row + 1
    for seq, item in enumerate(items, start=1):
        for idx, (_label, field) in enumerate(labels, start=1):
            if field == "_image":
                value = None
            elif field == "sort":
                value = seq
            else:
                value = getattr(item, field, None)
            cell = ws.cell(row=current, column=idx, value=value)
            cell.font = NORMAL_FONT
            cell.border = BORDER
            cell.alignment = WRAP
        if item.image_file_id and str(item.image_file_id) in images:
            _insert_image(ws, images[str(item.image_file_id)], f"B{current}", width=80)
            ws.row_dimensions[current].height = 62
        current += 1
    return current


def _fmt_date(value: Any) -> str:
    if not value:
        return ""
    try:
        return value.strftime("%Y-%m-%d")
    except AttributeError:
        return str(value)


# ------------------------------------------------------------------ 报价单


async def build_quotation_workbook(db: AsyncSession, quotation: Quotation, *, seller_default: dict[str, str]) -> bytes:
    items = list(
        (
            await db.execute(
                select(QuotationItem)
                .where(QuotationItem.quotation_id == quotation.id, QuotationItem.deleted_at.is_(None))
                .order_by(QuotationItem.sort, QuotationItem.created_at)
            )
        )
        .scalars()
        .all()
    )
    images = await _images(db, [i.image_file_id for i in items])
    company = (await db.execute(select(Company).where(Company.id == quotation.company_id))).scalar_one_or_none()

    wb = Workbook()
    ws = wb.active
    ws.title = "QUOTATION"

    ws["A1"] = quotation.seller_name or seller_default.get("name", "")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Add: {quotation.seller_address or seller_default.get('address', '')}"
    ws["A2"].font = NORMAL_FONT
    ws["A3"] = quotation.seller_contact or seller_default.get("contact", "")
    ws["A3"].font = NORMAL_FONT

    ws["A5"] = f"To: {company.name if company else ''}"
    ws["A5"].font = NORMAL_FONT
    ws["A6"] = quotation.title or "Offer"
    ws["A6"].font = TITLE_FONT

    labels = [
        "Item No.",
        "Picture",
        "Description of Goods",
        "Price",
        "MOQ",
        "Power",
        "IP",
        "Lumen",
        "Color TEMP.",
        "Color Rendering",
        "Price Valid",
    ]
    units = ["", "", "", quotation.currency, "", "(W)", "", "(lm)", "", "(Ra)", ""]

    header_row = 8
    for idx, label in enumerate(labels, start=1):
        cell = ws.cell(row=header_row, column=idx, value=label)
        cell.font = HEAD_FONT
        cell.fill = GREY
        cell.border = BORDER
        cell.alignment = CENTER
        unit_cell = ws.cell(row=header_row + 1, column=idx, value=units[idx - 1])
        unit_cell.font = NOTE_FONT
        unit_cell.border = BORDER
        unit_cell.alignment = CENTER

    # M2.5-D4：对外报价 2 位，xlsx 单位价同时落 2 位数值与 0.00 数字格式。
    price_col = labels.index("Price") + 1
    current = header_row + 2
    for item in items:
        values = [
            item.item_no or item.product_name,
            None,
            item.description or item.spec or "",
            round(float(item.unit_price or 0), 2),
            item.moq,
            item.power,
            item.ip,
            item.lumen,
            item.color_temp,
            item.color_rendering,
            item.price_valid or _fmt_date(quotation.valid_until),
        ]
        for idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current, column=idx, value=value)
            cell.font = NORMAL_FONT
            cell.border = BORDER
            cell.alignment = WRAP
            if idx == price_col:
                cell.number_format = "0.00"
        if item.image_file_id and str(item.image_file_id) in images:
            _insert_image(ws, images[str(item.image_file_id)], f"B{current}", width=100)
            ws.row_dimensions[current].height = 78
        current += 1

    current += 1
    for term in quotation.terms or []:
        ws.cell(row=current, column=1, value=term).font = NOTE_FONT
        current += 1

    if quotation.notes_cn:
        current += 1
        for line in quotation.notes_cn.splitlines():
            ws.cell(row=current, column=1, value=line).font = NOTE_FONT
            current += 1
    if quotation.notes_en:
        current += 1
        for line in quotation.notes_en.splitlines():
            ws.cell(row=current, column=1, value=line).font = NOTE_FONT
            current += 1

    widths = [26, 14, 52, 10, 9, 10, 8, 11, 12, 14, 12]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DEFAULT_TERMS = [
    "▲ Remarks：Price validity: 30 days from the date of quotation. This quotation is based on "
    "current raw material costs. Final price will be reconfirmed at the time of order confirmation.",
    "▲ If qty less than 1000pcs, extra start-up costs ￥800/ model",
    "▲ Color customized costs ￥2000/ model",
    "▲ FOB NINGBO",
    "▲ Payment term: 30% T/T in advance and the balance 70% should be settled on sight of B/L copy",
    "▲ Warranty: Two Years",
]

DEFAULT_NOTES_CN = """1.本报价基于以下标准配置:标准包装(彩盒+A4黑白说明书+黑字黄包)、标准起订量(1000个)、
无特殊认证、无定制要求。
2.如有特殊需求(彩盒、贴牌、认证、低于起订量、定制规格等)，需另行报价。
3.本报价有效期:30天(自报价之日起)。
4.正式订单下达前，请双方再次确认所有需求与报价单一致。如有变更，需重新报价。"""

DEFAULT_NOTES_EN = """Pricing basis: The above prices are for the basic color box + yellow outer carton with
black printed text + A4 black-and-white instruction manual.
For any special packaging requirements, such as hardening or thickening, or custom requests like
product logo silk-screen printing, please contact us separately to discuss the price difference."""
