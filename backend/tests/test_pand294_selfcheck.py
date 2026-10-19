"""PAND-294 自检断言：M2.5-D4 对外报价 2 位落实到导出产物。

场景：``final_price=148.20``、数量 7 → 单位价 ``21.1714…``，对外必须为 ``21.17``；
报价单 PDF（HTML + reportlab 两条路径）与 xlsx 单位价均为 2 位，不得出现 4 位对外价。
"""

from __future__ import annotations

import io
import uuid
from types import SimpleNamespace
from typing import Any, cast

import pytest
from openpyxl import load_workbook

from app.api.v1.endpoints.quotations import _allocate_unit_price
from app.services import export_excel
from app.services import pdf as pdf_service

FINAL_PRICE = 148.20
QUANTITY = 7
EXPECTED_UNIT_PRICE = 21.17


def _quotation_context(unit_price: float) -> dict:
    return pdf_service.build_context(
        seller={"name": "测试卖方", "address": "宁波", "contact": "sales@example.com"},
        quotation={
            "quotation_no": "QT20261019001",
            "created_at": "2026-10-19",
            "valid_until": "2026-11-18",
            "currency": "USD",
            "subtotal": round(QUANTITY * EXPECTED_UNIT_PRICE, 2),
            "additional_fee": 0,
            "discount": 0,
            "tax_amount": 0,
            "total_amount": round(QUANTITY * EXPECTED_UNIT_PRICE, 2),
        },
        customer={"name": "测试客户", "country": "US", "contact_name": "Buyer"},
        items=[
            {
                "product_name": "LED Panel",
                "spec": "600x600",
                "quantity": QUANTITY,
                "unit": "PCS",
                "unit_price": unit_price,
                "amount": round(QUANTITY * EXPECTED_UNIT_PRICE, 2),
            }
        ],
    )


class TestUnitPriceRounding:
    def test_from_pricing_allocates_two_decimals(self) -> None:
        # final_price=148.20 / 7 = 21.1714…，对外报价必须 2 位
        assert _allocate_unit_price(FINAL_PRICE, QUANTITY) == EXPECTED_UNIT_PRICE

    def test_margin_over_floor_still_two_decimals(self) -> None:
        assert _allocate_unit_price(FINAL_PRICE, QUANTITY, 0.02) == 21.59

    def test_fmt_renders_two_decimals(self) -> None:
        assert pdf_service._fmt(21.1714) == "21.17"


class TestQuotationPdf:
    def test_html_path_two_decimals(self) -> None:
        html = pdf_service.render_html(_quotation_context(21.1714))
        assert "21.17" in html
        assert "21.1714" not in html

    def test_reportlab_path_two_decimals(self) -> None:
        pdfminer = pytest.importorskip("pdfminer.high_level")
        data = pdf_service._reportlab_pdf(_quotation_context(21.1714))
        text = pdfminer.extract_text(io.BytesIO(data))
        assert "21.17" in text
        assert "21.1714" not in text

    def test_price_list_html_two_decimals(self) -> None:
        html = pdf_service.render_price_list_html(
            {
                "currency": "USD",
                "items": [{"item_no": "LED-01", "description": "LED Panel", "unit_price": 21.1714}],
            }
        )
        assert "21.17" in html
        assert "21.1714" not in html


class _Result:
    def __init__(self, rows: list) -> None:
        self._rows = rows

    def scalars(self) -> _Result:
        return self

    def all(self) -> list:
        return self._rows

    def scalar_one_or_none(self):
        return self._rows[0] if self._rows else None


class _FakeDb:
    """按调用顺序返回：明细 → 客户（有图片时才多一次文件查询）。"""

    def __init__(self, items: list, company) -> None:
        self._results = [items, [company]]

    async def execute(self, _stmt):
        return _Result(self._results.pop(0))


class TestQuotationExcel:
    async def test_xlsx_unit_price_two_decimals(self) -> None:
        item = SimpleNamespace(
            item_no="LED-01",
            product_name="LED Panel",
            description="LED Panel 600x600",
            spec="600x600",
            unit_price=21.1714,
            moq=1000,
            power="36W",
            ip="IP44",
            lumen="3600",
            color_temp="4000K",
            color_rendering="Ra80",
            price_valid="2026-11-18",
            image_file_id=None,
        )
        quotation = SimpleNamespace(
            id=uuid.uuid4(),
            company_id=uuid.uuid4(),
            seller_name="测试卖方",
            seller_address="宁波",
            seller_contact="sales@example.com",
            title="Offer",
            currency="USD",
            valid_until=SimpleNamespace(isoformat=lambda: "2026-11-18"),
            terms=["FOB NINGBO"],
            notes_cn=None,
            notes_en=None,
        )
        company = SimpleNamespace(name="测试客户")

        data = await export_excel.build_quotation_workbook(
            cast(Any, _FakeDb([item], company)), cast(Any, quotation), seller_default={}
        )
        ws = load_workbook(io.BytesIO(data)).active

        # 表头第 8 行，数据自第 10 行起；Price 为第 4 列。
        price_cell = ws.cell(row=10, column=4)
        assert price_cell.value == EXPECTED_UNIT_PRICE
        assert price_cell.number_format == "0.00"
